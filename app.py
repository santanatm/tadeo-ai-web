import streamlit as st
from groq import Groq
from langchain_community.tools.tavily_search import TavilySearchResults
from langchain_community.utilities.tavily_search import TavilySearchAPIWrapper
from fpdf import FPDF
import PyPDF2
from docx import Document
import io

# 1. PAGE CONFIGURATION
st.set_page_config(page_title="Tadeo AI", page_icon="🤖", layout="centered")

st.markdown("""
    <style>
    #MainMenu {visibility: hidden;}
    footer {display: none !important;}
    header {display: none !important;}
    .block-container {padding-top: 1rem !important;}
    .stFileUploader section { border: 1px solid #bfa34b !important; border-radius: 5px; }
    .stDownloadButton > button {
        background-color: #333333 !important;
        color: #bfa34b !important;
        border: 1px solid #bfa34b !important;
        width: 100%;
        font-weight: bold;
    }
    </style>
    """, unsafe_allow_html=True)

# 2. INTERFACE: PHOTO AND TITLE
col_foto, col_titulo = st.columns([1, 5]) 
with col_foto:
    try:
        st.image("Tadeo_Santana.png", width=60) 
    except:
        st.write("🤖")
with col_titulo:
    st.markdown("<h1 style='margin-top: -10px;'>TADEO AI</h1>", unsafe_allow_html=True)

# 3. KEYS AND ENGINES
GROQ_KEY = st.secrets.get("GROQ_API_KEY")
TAVILY_KEY = st.secrets.get("TAVILY_API_KEY")

if not GROQ_KEY or not TAVILY_KEY:
    st.error("API Keys missing in Secrets.")
    st.stop()

try:
    client = Groq(api_key=GROQ_KEY)
    api_wrapper = TavilySearchAPIWrapper(tavily_api_key=TAVILY_KEY)
    search = TavilySearchResults(api_wrapper=api_wrapper)
except Exception as e:
    st.error(f"Initialization error: {e}")

# 4. DOCUMENT TEXT EXTRACTION
def get_file_content(file):
    if file.type == "application/pdf":
        reader = PyPDF2.PdfReader(file)
        return " ".join([page.extract_text() for page in reader.pages])
    elif file.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        doc = Document(file)
        return " ".join([para.text for para in doc.paragraphs])
    else: # Plain text
        return file.read().decode("utf-8")

def crear_pdf(texto, consulta):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", 'B', 16)
    pdf.cell(200, 10, txt="Tadeo AI Report", ln=True, align='C')
    pdf.ln(10)
    pdf.set_font("Arial", 'B', 12)
    pdf.multi_cell(0, 10, txt=f"Analysis/Question: {consulta}")
    pdf.ln(5)
    pdf.set_font("Arial", size=11)
    texto_seguro = texto.encode('latin-1', 'replace').decode('latin-1')
    pdf.multi_cell(0, 10, txt=texto_seguro)
    return pdf.output(dest='S').encode('latin-1')

# 5. INPUT SECTION
archivo_subido = st.file_uploader("Upload a file to analyze", type=['pdf', 'txt', 'docx'])
pregunta = st.text_input("", placeholder="What do you want to know? (e.g. Who was Albert Einstein?)", key="user_input")

# 6. AUTOMATIC EXECUTION LOGIC
if pregunta or archivo_subido:
    with st.spinner("Tadeo is processing your request..."):
        try:
            full_context = ""
            label_name = pregunta if pregunta else f"Analysis of {archivo_subido.name}"

            # If a file is uploaded, extract its actual text
            if archivo_subido:
                file_text = get_file_content(archivo_subido)
                full_context += f"\n[Document Content: {file_text}]\n"

            # Web Search
            try:
                search_query = pregunta if pregunta else f"Summary of {archivo_subido.name}"
                resultados_web = search.run(search_query)
                if "401" not in str(resultados_web):
                    full_context += f"\n[Web Context: {resultados_web}]\n"
            except:
                pass

            # AI Generation
            chat_completion = client.chat.completions.create(
                messages=[
                    {"role": "system", "content": "You are Tadeo AI. Respond in English. Use the provided document content and web context to give a detailed answer. Ignore API technical errors."},
                    {"role": "user", "content": f"Context: {full_context}\n\nUser Question/Task: {label_name}"}
                ],
                model="llama-3.3-70b-versatile",
            )
            
            respuesta_final = chat_completion.choices[0].message.content
            st.subheader("📝 Result:")
            st.write(respuesta_final)
            
            # PDF Download
            pdf_bytes = crear_pdf(respuesta_final, label_name)
            st.download_button(label="📥 Download report as PDF", data=pdf_bytes, file_name="Tadeo_AI_Report.pdf", mime="application/pdf")
            
        except Exception as e:
            st.error(f"An error occurred: {e}")
